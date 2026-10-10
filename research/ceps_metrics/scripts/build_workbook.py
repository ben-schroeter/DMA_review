"""Assemble research/ceps_metrics/ceps_instruments_metrics.xlsx from data/*.json|csv."""
import csv, json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

D = "../data/"
rows = list(csv.DictReader(open(D + "ceps_instruments.csv")))
imap = list(csv.DictReader(open(D + "instrument_map.csv")))
meas = {(m["celex"], m["kind"]): m for m in json.load(open(D + "measurements.json"))}
mainm = {m["celex"]: m for m in json.load(open(D + "measurements.json")) if m["kind"] in ("main", "secondary")}
cons = json.load(open(D + "consolidated.json"))
nim = json.load(open(D + "national_measures.json"))
sec = json.load(open(D + "secondary_acts.json"))
web = list(csv.DictReader(open(D + "guidance_web_supplement.csv")))
val = json.load(open(D + "validation_results.json"))
pol = json.load(open(D + "policy_area.json"))
import os as _os0
_imp_all = {json.load(open(D + "impact/" + f))["celex"]: json.load(open(D + "impact/" + f))
            for f in (_os0.listdir(D + "impact") if _os0.path.isdir(D + "impact") else []) if f.endswith(".json")}
import re as _re
ag = []
for f in ("agency_edpb.json", "agency_eba.json", "agency_esma.json"):
    for it in json.load(open(D + f)):
        if it.get("parent") and it.get("final", True):
            ag.append(it)
_seen = {}
def _norm(t):
    return _re.sub(r"[^a-z]", "", (t or "").lower().replace("final report on", "").replace("final report", ""))[:45]
dedup = []
_words = lambda t: set(w for w in _re.findall(r"[a-z]{4,}", (t or "").lower()) if w not in ("final", "report", "guidelines", "joint"))
for it in ag:
    def _same(a, b):
        A, B = _words(a.get("title")), _words(b.get("title"))
        if a.get("agency") == b.get("agency"):
            return _norm(a.get("title")) == _norm(b.get("title"))
        return min(len(A), len(B)) >= 2 and len(A & B) / min(len(A), len(B)) >= 0.8
    dup = next((d for d in dedup if d["parent"] == it["parent"] and _same(d, it)), None)
    if dup:
        if dup.get("agency") != it.get("agency"):
            dup["joint_also_published_by"] = it.get("agency")
        continue
    dedup.append(it)
ag = dedup
da = json.load(open(D + "dates_amendments.json"))
props = json.load(open(D + "proposals.json"))
NOT_COMPARABLE = {"32006L0112": "recast proposal: PDF includes struck-through and correlation material",
                  "32002R0006": "1993 proposal PDF bundles the companion Design Directive proposal",
                  "32004R0139": "proposal COM(2002) 711 not available in EUR-Lex"}

URL = "https://eur-lex.europa.eu/legal-content/EN/TXT/?uri=CELEX:{}"
METRICS = ["pages", "words_legal_text", "words_preamble", "words_recitals", "words_articles", "words_annexes",
           "words_memorandum", "words_financial_statement", "restrictions_total", "r_obligation", "r_prohibition",
           "r_restricted_permission", "r_explicit_requirement", "r_duty_of_result", "addr_private_parties",
           "addr_member_states", "addr_eu_institutions", "addr_unclear"]
HDR = Font(name="Arial", bold=True, color="FFFFFF"); FILL = PatternFill("solid", fgColor="1F4E78")
BODY = Font(name="Arial", size=10)


def style(ws, widths):
    for c in ws[1]:
        c.font = HDR; c.fill = FILL; c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=10, color=c.font.color.rgb if c.font and c.font.color and c.font.color.type == "rgb" else None)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 45


def summed(celexes, kind="main"):
    out, notes = {}, []
    for c in celexes:
        m = meas.get((c, kind)) or mainm.get(c)
        if not m or "error" in m:
            notes.append(f"{c}: not measured ({(m or {}).get('error', 'missing')})"); continue
        if m.get("parse_warning"):
            notes.append(f"{c}: {m['parse_warning']}")
        for k in METRICS:
            v = m.get(k)
            if v is not None:
                out[k] = out.get(k, 0) + v
    return out, notes


wb = Workbook()
# ---------------- secondary sheet first (instrument formulas reference it) ----------------
ws2 = wb.active; ws2.title = "secondary_acts"
row_of = {}
for r in imap:
    for c in r["celex_measured"].split(";"):
        row_of.setdefault(c, (r["row_id"], r["name"]))
sh = ["parent_row_id", "parent_name", "parent_celex", "celex", "kind", "eurlex_type", "date", "title",
      "addressed_to_specific_member_state", "source_url"] + METRICS + ["note"]
ws2.append(sh)
for s in sorted(sec, key=lambda s: (int(row_of.get(s["parent"], ("999",))[0]), s["date"], s["celex"])):
    rid, nm = row_of.get(s["parent"], ("", ""))
    m, notes = summed([s["celex"]], "secondary")
    note = "; ".join(notes)
    if s["kind"] == "guidance":
        note = (note + "; " if note else "") + "Guidance is non-binding: restriction counts mostly reflect obligations it restates"
    ws2.append([int(rid) if rid else None, nm, s["parent"], s["celex"], s["kind"], s["type"], s["date"], s["title"],
                "yes" if s["addressed_to_specific_state_or_party"] else "no", URL.format(s["celex"])]
               + [m.get(k) for k in METRICS] + [note])
n2 = ws2.max_row
style(ws2, [8, 30, 13, 16, 14, 12, 11, 60, 12, 30] + [10] * len(METRICS) + [40])

# ---------------- web guidance ----------------
ws3 = wb.create_sheet("guidance_web")
ws3.append(["parent_row_id", "parent_name", "parent_celex", "reference", "date", "title", "source_url", "note"])
for g in web:
    rid, nm = row_of.get(g["parent_celex"], ("", ""))
    ws3.append([int(rid) if rid else None, nm, g["parent_celex"], g["reference"], g["date"], g["title"], g["source_url"],
                g["note"] + "; found by web search, not linked in EUR-Lex; not measured"])
n3 = ws3.max_row
style(ws3, [8, 30, 13, 14, 11, 70, 50, 50])

# ---------------- agency guidance ----------------
wg = wb.create_sheet("guidance_agencies")
gcols = ["parent_row_id", "parent_name", "parent_celex", "agency", "title", "date", "status", "pdf_url", "pages", "words_total",
         "restrictions_total", "r_obligation", "r_prohibition", "r_restricted_permission", "r_explicit_requirement",
         "r_duty_of_result", "addr_private_parties", "addr_member_states", "addr_eu_institutions", "addr_unclear",
         "should_count", "note"]
wg.append(gcols)
for it in sorted(ag, key=lambda x: (int(row_of.get(x["parent"], ("999",))[0]), x.get("agency", ""), str(x.get("date") or x.get("date_text") or ""))):
    rid, nm = row_of.get(it["parent"], ("", ""))
    notes = []
    if it.get("endorsed_wp29"): notes.append("Article 29 Working Party guidelines endorsed by the EDPB")
    if it.get("joint_also_published_by"): notes.append(f"joint guidelines, also published by {it['joint_also_published_by']}")
    if it.get("also_relates_to"): notes.append(f"also relates to {it['also_relates_to']}")
    if it.get("agency") in ("EBA", "ESMA"): notes.append("measured on the final report (includes rationale and consultation feedback)")
    if not it.get("pages"): notes.append("not measured: " + (it.get("error") or "PDF on a domain not reachable from this environment"))
    notes.append("non-binding guidance; 'should' is its usual verb (should_count)")
    wg.append([int(rid) if rid else None, nm, it["parent"], it.get("agency"), it.get("title"), it.get("date") or it.get("date_text"),
               "final", it.get("pdf") or it.get("url"), it.get("pages"), it.get("words_total")] +
              [it.get(k) for k in gcols[10:21]] + ["; ".join(notes)])
ng = wg.max_row
style(wg, [8, 30, 13, 8, 70, 11, 8, 40] + [10] * 13 + [60])

import re as _re_c
CAUSAL = ("causal", "difference-in-difference", "event study", "event-study", "synthetic control", "regression discontinuity",
          "natural experiment", "instrumental variable", "randomi")


def is_causal(m):
    did = bool(_re_c.search(r"\bDiD\b|\bDID\b", m))
    m = m.lower()
    return (did or any(k in m for k in CAUSAL)) and "non-causal" not in m and "not causal" not in m


# ---------------- instruments ----------------
ws = wb.create_sheet("instruments", 0)
cols = ["row_id", "name", "identifier", "policy_area_ceps", "ceps_table_status", "proposal_date", "confidence_proposal_date",
        "adoption_date", "oj_publication_date", "entry_into_force", "status_ceps_jul2025", "status_oct2026",
        "celex_measured", "text_measured", "source_url"] + METRICS + \
       ["consolidated_celex", "consolidated_restrictions_total", "consolidated_addr_private_parties",
        "n_implementing_acts", "n_delegated_acts", "n_other_commission_council_acts", "n_guidance_eurlex", "n_guidance_web",
        "n_guidance_edpb", "n_guidance_eba", "n_guidance_esma",
        "restrictions_in_implementing_and_delegated_acts", "national_transposition_measures", "member_states_with_measures",
        "n_amending_acts", "first_amendment_date", "latest_amendment_date", "n_corrigenda",
        "proposal_celex", "proposal_text_source", "proposal_words_legal_text", "proposal_restrictions_total",
        "proposal_addr_private_parties", "change_words", "change_words_pct", "change_restrictions", "change_restrictions_pct",
        "change_restrictions_private", "impact_evidence_level", "impact_evidence_rating", "impact_summary", "n_promised_indicators", "n_indicators_measured",
        "n_impact_studies", "n_causal_studies", "note"]
ws.append(cols)
ci = {c: i + 1 for i, c in enumerate(cols)}
for i, (r, m0) in enumerate(zip(rows, imap), start=2):
    cel = m0["celex_measured"].split(";")
    if cel[0].startswith("COUNCIL"):
        text, url = "Council final compromise text, 23 Apr 2026", ""
        met, notes = {}, ["Not measured: Council document register (data.consilium.europa.eu) is blocked by this environment's network policy"]
    else:
        text = "Commission proposal (COM)" if cel[0].startswith("5") else "Original act as published in the OJ"
        url = " ; ".join(URL.format(c) for c in cel)
        met, notes = summed(cel)
    cc = [cons[c]["latest"] for c in cel if c in cons]
    cmet, cnotes = {}, []
    if cc:
        ok = [x for x in cc if "error" not in meas.get((x, "consolidated"), {"error": 1})]
        if ok:
            cmet, cnotes = summed(ok, "consolidated")
        else:                           # EUR-Lex lists only the initial version: never amended
            cmet = {k: met.get(k) for k in ("restrictions_total", "addr_private_parties")}
            cnotes = ["no consolidated text with amendments in EUR-Lex; consolidated count = original"]
    dirs = [c for c in cel if c[5:6] == "L"]
    if "Block Exemption" in r["name"]:
        notes.append("Block exemption: defines a safe harbour ('the exemption shall not apply to...'), so few direct restrictions are counted")
    note = "; ".join(x for x in [m0["note"]] + notes + cnotes if x)
    pa = pol[r["name"]]
    dd = [da["dates"].get(c, {}) for c in cel]
    first = lambda k: min((d.get(k) for d in dd if d.get(k)), default=None)
    vals = [int(m0["row_id"]), r["name"], r["identifier"], pa["policy_area"], pa["ceps_colour_status"], r["proposal_date"], r["confidence"],
            first("adoption_date"), first("oj_publication_date"), first("entry_into_force"), m0["status_ceps_jul2025"],
            m0["status_oct2026"], ";".join(cel), text, url] + [met.get(k) for k in METRICS] + \
           [";".join(cc) or None, cmet.get("restrictions_total"), cmet.get("addr_private_parties")]
    _id = next((_imp_all[c] for c in cel if c in _imp_all), None)
    if _id:
        _ind = [x for x in _id["indicators"] if x["indicator"]]
        _ev = [x for x in _id["literature"] if x.get("category", "evidence") == "evidence"]
        imp_vals = [_id.get("evidence_level"), _id["evidence_rating"], _id["summary"], len(_ind),
                    sum(1 for x in _ind if x.get("status_group", "not measured") != "not measured"),
                    sum(1 for x in _ev), sum(1 for x in _ev if is_causal(x["method"]))]
    else:
        imp_vals = [None] * 7
    am = [a for a in da["amendments"] if a["parent"] in cel and a["kind"] == "amendment"]
    co = [a for a in da["amendments"] if a["parent"] in cel and a["kind"] == "corrigendum"]
    amd = sorted(a["date"] for a in am if a["date"])
    adopted = cel[0][:1] == "3"
    # proposal comparison (adopted acts only)
    pc_vals = [None] * 5
    if adopted:
        pcs = [props["links"].get(c, {}).get("proposal_celex") for c in cel]
        why = next((NOT_COMPARABLE[c] for c in cel if c in NOT_COMPARABLE), None)
        if all(pcs) and not why:
            pm = [props["measurements"][p] for p in pcs]
            src = "PDF text (approximate)" if any(p.get("source") for p in pm) else "EUR-Lex HTML"
            pc_vals = [";".join(pcs), src, sum(p["words_legal_text"] for p in pm), sum(p["restrictions_total"] for p in pm),
                       sum(p["addr_private_parties"] for p in pm)]
        elif why:
            pc_vals = [";".join(p for p in pcs if p) or None, "not comparable: " + why, None, None, None]
        else:
            pc_vals = [None, "no legislative proposal (Commission act)" if "Block Exemption" in r["name"] else "no linked proposal", None, None, None]
    else:
        pc_vals = [None, "row is itself a proposal", None, None, None]
    ws.append(vals + [None] * 9 + [sum(nim[c]["n_measures"] for c in dirs if c in nim) if dirs else None,
                                   max((nim[c]["n_member_states"] for c in dirs if c in nim), default=None) if dirs else None,
                                   len(am) if adopted else None, amd[0] if amd else None, amd[-1] if amd else None,
                                   len(co) if adopted else None] + pc_vals + [None] * 5 + imp_vals + [note])
    L = lambda k: get_column_letter(ci[k])
    pw, aw, pr, ar = f"{L('proposal_words_legal_text')}{i}", f"{L('words_legal_text')}{i}", f"{L('proposal_restrictions_total')}{i}", f"{L('restrictions_total')}{i}"
    pp, ap = f"{L('proposal_addr_private_parties')}{i}", f"{L('addr_private_parties')}{i}"
    ws.cell(row=i, column=ci["change_words"], value=f'=IF(ISNUMBER({pw}),{aw}-{pw},"")')
    ws.cell(row=i, column=ci["change_words_pct"], value=f'=IF(AND(ISNUMBER({pw}),{pw}>0),({aw}-{pw})/{pw},"")')
    ws.cell(row=i, column=ci["change_restrictions"], value=f'=IF(ISNUMBER({pr}),{ar}-{pr},"")')
    ws.cell(row=i, column=ci["change_restrictions_pct"], value=f'=IF(AND(ISNUMBER({pr}),{pr}>0),({ar}-{pr})/{pr},"")')
    ws.cell(row=i, column=ci["change_restrictions_private"], value=f'=IF(ISNUMBER({pp}),{ap}-{pp},"")')
    for k in ("change_words_pct", "change_restrictions_pct"):
        ws.cell(row=i, column=ci[k]).number_format = "0.0%"
    R = f"secondary_acts!$A$2:$A${n2}"; K = f"secondary_acts!$E$2:$E${n2}"
    RT = f"secondary_acts!${get_column_letter(10 + METRICS.index('restrictions_total') + 1)}$2:${get_column_letter(10 + METRICS.index('restrictions_total') + 1)}${n2}"
    f = {"n_implementing_acts": f'=COUNTIFS({R},$A{i},{K},"implementing")',
         "n_delegated_acts": f'=COUNTIFS({R},$A{i},{K},"delegated")',
         "n_other_commission_council_acts": f'=COUNTIFS({R},$A{i},{K},"other Commission/Council act")',
         "n_guidance_eurlex": f'=COUNTIFS({R},$A{i},{K},"guidance")',
         "n_guidance_web": f'=COUNTIF(guidance_web!$A$2:$A${max(n3, 2)},$A{i})',
         "n_guidance_edpb": f'=COUNTIFS(guidance_agencies!$A$2:$A${ng},$A{i},guidance_agencies!$D$2:$D${ng},"EDPB")',
         "n_guidance_eba": f'=COUNTIFS(guidance_agencies!$A$2:$A${ng},$A{i},guidance_agencies!$D$2:$D${ng},"EBA")',
         "n_guidance_esma": f'=COUNTIFS(guidance_agencies!$A$2:$A${ng},$A{i},guidance_agencies!$D$2:$D${ng},"ESMA")',
         "restrictions_in_implementing_and_delegated_acts":
             f'=SUMIFS({RT},{R},$A{i},{K},"implementing")+SUMIFS({RT},{R},$A{i},{K},"delegated")'}
    for k, v in f.items():
        ws.cell(row=i, column=ci[k], value=v)
# totals row
tr = ws.max_row + 1
ws.cell(row=tr, column=2, value="TOTAL (all rows)").font = Font(name="Arial", bold=True)
for k in METRICS + ["consolidated_restrictions_total", "consolidated_addr_private_parties", "n_implementing_acts",
                    "n_delegated_acts", "n_other_commission_council_acts", "n_guidance_eurlex", "n_guidance_web",
                    "n_guidance_edpb", "n_guidance_eba", "n_guidance_esma", "restrictions_in_implementing_and_delegated_acts", "national_transposition_measures",
                    "n_amending_acts", "n_corrigenda"]:
    Lc = get_column_letter(ci[k]); ws.cell(row=tr, column=ci[k], value=f"=SUM({Lc}2:{Lc}{tr - 1})")
style(ws, [6, 34, 22, 18, 12, 11, 9, 11, 11, 11, 10, 10, 16, 18, 30] + [10] * len(METRICS) + [20, 11, 11, 9, 9, 9, 9, 9, 9, 9, 9, 12, 11, 10]
      + [9, 11, 11, 9, 16, 18, 10, 10, 10, 10, 9, 10, 9, 10, 30, 50, 9, 9, 9, 9, 60])
for c in ws[tr]:
    c.font = Font(name="Arial", bold=True, size=10)
for k in METRICS + ["consolidated_restrictions_total", "restrictions_in_implementing_and_delegated_acts", "national_transposition_measures"]:
    for row in ws.iter_rows(min_row=2, min_col=ci[k], max_col=ci[k]):
        row[0].number_format = "#,##0"

# ---------------- amendments ----------------
wa = wb.create_sheet("amendments", 2)
wa.append(["parent_row_id", "parent_name", "parent_celex", "celex", "kind", "date", "title", "source_url"])
for a in sorted(da["amendments"], key=lambda a: (int(row_of.get(a["parent"], ("999",))[0]), a["kind"], a["date"])):
    rid, nm = row_of.get(a["parent"], ("", ""))
    wa.append([int(rid) if rid else None, nm, a["parent"], a["celex"], a["kind"], a["date"], a["title"], URL.format(a["celex"])])
style(wa, [8, 30, 13, 18, 12, 11, 80, 40])

# ---------------- external benchmarks ----------------
import os as _os
wbm = wb.create_sheet("benchmarks")
wbm.append(["benchmark", "variant", "pages", "words", "restrictions_total", "obligation", "prohibition", "explicit_requirement",
            "restrictions_per_1000_words", "private_parties", "US_States", "federal_government_regulators", "unclear",
            "validation_precision", "precision_adjusted_restrictions", "note"])
for f in sorted(x for x in _os.listdir(D) if x.startswith("benchmark_") and x.endswith(".json") and "_instances" not in x):
    bm = json.load(open(D + f))
    for var in ("all_text", "excluding_inserted_text"):
        v = bm[var]; a = v["by_addressee"]; c = v["by_category"]; prec = bm.get("validation", {}).get("precision")
        rr = wbm.max_row + 1
        wbm.append([bm["label"].upper(), var.replace("_", " "), bm["pages"], v["words"], v["restrictions_total"], c.get("obligation"),
                    c.get("prohibition"), c.get("explicit_requirement"), v["per_1000_words"],
                    a.get("private parties"), a.get("US States"), a.get("federal government / regulators (SEC, PCAOB)"), a.get("unclear / none"),
                    prec, f"=E{rr}*N{rr}" if prec else None,
                    (bm.get("source", "") + ". " + bm.get("validation", {}).get("note", "")) if var == "all_text" else
                    "Excludes text inserted into other statutes (in quotation marks), mirroring how quoted amending text is skipped for EU acts"])
wbm.append([])
wbm.append(["EU comparison (105 adopted acts, original texts)", "mean per act", None, 14601, 207, None, None, None, 14.0, None, None, None, None, 0.84,
            "=E" + str(wbm.max_row + 1) + "*N" + str(wbm.max_row + 1), "Words = articles + annexes (no recitals), comparable to a US statute's text"])
style(wbm, [22, 24, 8, 10, 12, 10, 10, 10, 12, 10, 10, 14, 10, 12, 14, 80])
wbm.freeze_panes = None

# ---------------- impact evidence (pilot) ----------------
_imp = {}
_idir = D + "impact/"
if _os.path.isdir(_idir):
    for f in sorted(_os.listdir(_idir)):
        if f.endswith(".json"):
            _d = json.load(open(_idir + f)); _imp[_d["celex"]] = _d
wpi = wb.create_sheet("proposal_indicators")
wpi.append(["parent_row_id", "parent_name", "proposal_celex", "type", "objective", "indicator / promised impact", "target or baseline",
            "evidence found", "status", "status_group", "sources"])
wil = wb.create_sheet("impact_literature")
wil.append(["parent_row_id", "parent_name", "category", "id", "citation", "year", "source_type", "method", "outcomes", "direction", "headline finding",
            "plain-language summary", "period / geography", "quality notes", "interested party", "access", "verified", "accessed", "url",
            "pdf_url", "archive_url"])
for c, d_ in _imp.items():
    rid, nm = row_of.get(c, ("", ""))
    for x in d_["indicators"]:
        wpi.append([int(rid) if rid else None, nm, d_["proposal"], x["type"], x["objective"], x["indicator"], x["target_or_baseline"],
                    x["evidence"], x["status"], x.get("status_group"), x["sources"]])
    for x in d_["literature"]:
        wil.append([int(rid) if rid else None, nm, x.get("category"), x.get("id"), x["citation"], x["year"], x["source_type"], x["method"], x["outcomes"],
                    x["direction"], x["headline"], x.get("summary_plain"), x["period_geography"], x["quality"], x.get("interested_party"),
                    x.get("access"), x["verified"], x.get("accessed"), x["url"], x.get("pdf_url"), x.get("archive_url")])
for w_, widths in ((wpi, [8, 22, 13, 18, 40, 45, 14, 70, 22, 16, 30]), (wil, [8, 22, 16, 22, 60, 6, 18, 30, 25, 18, 60, 60, 18, 40, 30, 14, 14, 11, 40, 30, 40])):
    style(w_, widths)
    for row in w_.iter_rows(min_row=2):
        for c_ in row:
            c_.alignment = Alignment(wrap_text=True, vertical="top")

# ---------------- method + validation ----------------
wm = wb.create_sheet("method")
for line in open(D + "../METHOD.md").read().splitlines():
    wm.append([line])
wm.column_dimensions["A"].width = 140
for row in wm.iter_rows():
    row[0].font = Font(name="Arial", size=10, bold=row[0].value is not None and str(row[0].value).startswith("#"))
    row[0].alignment = Alignment(wrap_text=True)
wv = wb.create_sheet("validation")
wv.append(["sample", "counted_true", "counted_false", "missed_true", "correctly_dropped", "precision", "recall_within_sample",
           "addressee_accuracy", "addressee_n"])
for k, v in val.items():
    if isinstance(v, dict) and "precision" in v:
        wv.append([k] + [v[x] for x in ["counted_true", "counted_false", "missed_true", "correctly_dropped", "precision",
                                         "recall_within_sample", "addressee_accuracy", "addressee_n"]])
    elif isinstance(v, dict):
        wv.append([k] + [json.dumps(v)])
    else:
        wv.append([k, str(v)])
style(wv, [28, 12, 12, 12, 14, 10, 14, 14, 12])
wv.freeze_panes = None
wb.save("../ceps_instruments_metrics.xlsx")
print("saved", ws.max_row, n2, n3)
